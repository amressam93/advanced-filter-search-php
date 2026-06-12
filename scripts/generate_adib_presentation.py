from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


BASE_DIR = Path(__file__).resolve().parent.parent
ASSETS_DIR = BASE_DIR / "assets" / "adib"
OUTPUT_DIR = BASE_DIR / "presentations"
OUTPUT_FILE = OUTPUT_DIR / "adib_simple_visual_presentation.pptx"


def add_image_contain(slide, image_path, left, top, width, height):
    """Add image inside a bounding box while preserving aspect ratio."""
    with Image.open(image_path) as img:
        img_w, img_h = img.size

    box_ratio = width / height
    img_ratio = img_w / img_h

    if img_ratio > box_ratio:
        draw_w = width
        draw_h = width / img_ratio
        draw_left = left
        draw_top = top + (height - draw_h) / 2
    else:
        draw_h = height
        draw_w = height * img_ratio
        draw_top = top
        draw_left = left + (width - draw_w) / 2

    slide.shapes.add_picture(str(image_path), draw_left, draw_top, draw_w, draw_h)


def add_card_frame(slide, left, top, width, height):
    frame = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    frame.fill.solid()
    frame.fill.fore_color.rgb = RGBColor(255, 255, 255)
    frame.line.color.rgb = RGBColor(222, 229, 239)
    frame.line.width = Pt(1.2)
    frame.shadow.inherit = False
    return frame


def set_text_style(text_frame, font_size, bold=False, color=RGBColor(31, 41, 55)):
    for p in text_frame.paragraphs:
        for run in p.runs:
            run.font.size = Pt(font_size)
            run.font.bold = bold
            run.font.color.rgb = color


def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank = prs.slide_layouts[6]

    logo_image = ASSETS_DIR / "logo_full_header.png"
    card_images = [
        ASSETS_DIR / "card1_cashback_platinum.png",
        ASSETS_DIR / "card2_exceed_platinum.png",
        ASSETS_DIR / "card3_dana_master.png",
    ]

    # Slide 1: requested layout (logo left, 3 images right)
    slide = prs.slides.add_slide(blank)
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = RGBColor(247, 250, 255)

    title_box = slide.shapes.add_textbox(Inches(0.75), Inches(0.35), Inches(6.2), Inches(0.7))
    title_tf = title_box.text_frame
    title_tf.text = "ADIB Card Presentation"
    set_text_style(title_tf, 34, bold=True, color=RGBColor(15, 43, 90))

    subtitle_box = slide.shapes.add_textbox(
        Inches(0.75), Inches(1.0), Inches(6.2), Inches(0.5)
    )
    subtitle_tf = subtitle_box.text_frame
    subtitle_tf.text = "Simple | Visual | Readable"
    set_text_style(subtitle_tf, 16, color=RGBColor(64, 85, 119))

    left_panel = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.65), Inches(1.55), Inches(5.8), Inches(5.35)
    )
    left_panel.fill.solid()
    left_panel.fill.fore_color.rgb = RGBColor(255, 255, 255)
    left_panel.line.color.rgb = RGBColor(222, 229, 239)
    left_panel.line.width = Pt(1.5)
    left_panel.shadow.inherit = False

    add_image_contain(
        slide, logo_image, Inches(1.0), Inches(2.0), Inches(5.1), Inches(4.4)
    )

    right_x = Inches(7.0)
    row_top = [Inches(1.45), Inches(3.1), Inches(4.75)]
    card_w = Inches(5.55)
    card_h = Inches(1.45)

    for i, image in enumerate(card_images):
        add_card_frame(slide, right_x, row_top[i], card_w, card_h)
        add_image_contain(
            slide,
            image,
            right_x + Inches(0.18),
            row_top[i] + Inches(0.08),
            card_w - Inches(0.36),
            card_h - Inches(0.16),
        )

    note_box = slide.shapes.add_textbox(
        Inches(7.05), Inches(6.35), Inches(5.45), Inches(0.45)
    )
    note_tf = note_box.text_frame
    note_tf.text = "Card visuals from ADIB public pages"
    set_text_style(note_tf, 12, color=RGBColor(106, 123, 148))

    # Slide 2: card lineup overview
    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)

    title = slide.shapes.add_textbox(Inches(0.7), Inches(0.45), Inches(12), Inches(0.6))
    title.text_frame.text = "Card Lineup at a Glance"
    set_text_style(title.text_frame, 30, bold=True, color=RGBColor(16, 44, 95))

    labels = ["Cashback Platinum", "Exceed Platinum", "Dana Master Card"]
    left_positions = [Inches(0.8), Inches(4.65), Inches(8.5)]
    for idx, img in enumerate(card_images):
        frame = add_card_frame(slide, left_positions[idx], Inches(1.35), Inches(3.95), Inches(3.35))
        frame.line.width = Pt(1.0)
        add_image_contain(
            slide, img, left_positions[idx] + Inches(0.22), Inches(1.65), Inches(3.5), Inches(2.1)
        )
        label = slide.shapes.add_textbox(left_positions[idx], Inches(5.0), Inches(3.95), Inches(0.5))
        label_tf = label.text_frame
        label_tf.text = labels[idx]
        label_tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        set_text_style(label_tf, 16, bold=True, color=RGBColor(25, 55, 108))

    # Slide 3: key benefits
    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(245, 248, 253)

    title = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(12), Inches(0.6))
    title.text_frame.text = "Key Benefits Snapshot"
    set_text_style(title.text_frame, 30, bold=True, color=RGBColor(16, 44, 95))

    info_panel = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.85), Inches(1.3), Inches(7.4), Inches(5.55)
    )
    info_panel.fill.solid()
    info_panel.fill.fore_color.rgb = RGBColor(255, 255, 255)
    info_panel.line.color.rgb = RGBColor(220, 228, 240)
    info_panel.line.width = Pt(1.3)

    bullets = [
        "Simple product tiers for different customer needs",
        "Clear branding with premium visual identity",
        "Strong focus on rewards and lifestyle value",
        "Readable design, easy to compare card options",
    ]
    bullet_box = slide.shapes.add_textbox(Inches(1.25), Inches(1.9), Inches(6.6), Inches(4.6))
    tf = bullet_box.text_frame
    tf.clear()
    for idx, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = f"• {bullet}"
        p.space_after = Pt(13)
        for run in p.runs:
            run.font.size = Pt(22)
            run.font.color.rgb = RGBColor(32, 50, 83)

    visual_panel = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(8.55), Inches(1.3), Inches(3.95), Inches(5.55)
    )
    visual_panel.fill.solid()
    visual_panel.fill.fore_color.rgb = RGBColor(255, 255, 255)
    visual_panel.line.color.rgb = RGBColor(220, 228, 240)
    visual_panel.line.width = Pt(1.3)
    add_image_contain(
        slide, card_images[0], Inches(8.8), Inches(1.95), Inches(3.45), Inches(1.45)
    )
    add_image_contain(
        slide, card_images[1], Inches(8.8), Inches(3.55), Inches(3.45), Inches(1.45)
    )
    add_image_contain(
        slide, card_images[2], Inches(8.8), Inches(5.15), Inches(3.45), Inches(1.45)
    )

    # Slide 4: recommendations
    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)

    title = slide.shapes.add_textbox(Inches(0.75), Inches(0.45), Inches(12), Inches(0.6))
    title.text_frame.text = "Simple Design Enhancements"
    set_text_style(title.text_frame, 30, bold=True, color=RGBColor(16, 44, 95))

    tips = [
        ("1", "Use one core brand color per slide"),
        ("2", "Keep text short: max 4 bullet lines"),
        ("3", "Use one main visual focal point"),
        ("4", "Maintain strong contrast for readability"),
    ]
    y = Inches(1.5)
    for number, text in tips:
        circle = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(1.0), y, Inches(0.5), Inches(0.5))
        circle.fill.solid()
        circle.fill.fore_color.rgb = RGBColor(25, 78, 161)
        circle.line.color.rgb = RGBColor(25, 78, 161)
        ctf = circle.text_frame
        ctf.text = number
        ctf.paragraphs[0].alignment = PP_ALIGN.CENTER
        for run in ctf.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(16)
            run.font.bold = True

        tip_box = slide.shapes.add_textbox(Inches(1.75), y - Inches(0.02), Inches(8.5), Inches(0.55))
        tip_tf = tip_box.text_frame
        tip_tf.text = text
        set_text_style(tip_tf, 22, color=RGBColor(34, 50, 78))
        y += Inches(1.2)

    add_image_contain(
        slide, logo_image, Inches(9.7), Inches(2.0), Inches(2.9), Inches(2.9)
    )

    # Slide 5: closing
    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(13, 36, 81)

    add_image_contain(
        slide, logo_image, Inches(5.1), Inches(1.15), Inches(3.2), Inches(2.4)
    )

    thanks = slide.shapes.add_textbox(Inches(3.2), Inches(4.15), Inches(6.9), Inches(0.9))
    thanks_tf = thanks.text_frame
    thanks_tf.text = "Thank You"
    thanks_tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    set_text_style(thanks_tf, 46, bold=True, color=RGBColor(255, 255, 255))

    summary = slide.shapes.add_textbox(Inches(2.2), Inches(5.3), Inches(9.2), Inches(0.8))
    summary_tf = summary.text_frame
    summary_tf.text = "A simple, visual, and readable ADIB presentation template."
    summary_tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    set_text_style(summary_tf, 18, color=RGBColor(206, 220, 246))

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT_FILE)
    print(f"Created: {OUTPUT_FILE}")


if __name__ == "__main__":
    build_presentation()
