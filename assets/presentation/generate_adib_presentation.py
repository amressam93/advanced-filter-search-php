from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


BASE_DIR = Path(__file__).resolve().parent
OUTPUT_FILE = Path("/workspace/ADIB_Simple_Visual_Presentation.pptx")

LOGO = BASE_DIR / "logo_site.png"
CARDS = {
    "Platinum Cash Back": BASE_DIR / "card_platinum_cashback.png",
    "Titanium Cash Back": BASE_DIR / "card_titanium_cashback.png",
    "Platinum Debit": BASE_DIR / "card_platinum_debit.png",
}

NAVY = RGBColor(7, 67, 138)
CYAN = RGBColor(47, 181, 234)
SKY = RGBColor(217, 244, 252)
SLATE = RGBColor(64, 72, 81)
LIGHT = RGBColor(245, 248, 251)
WHITE = RGBColor(255, 255, 255)


def set_background(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_textbox(slide, left, top, width, height, text, size, color, bold=False,
                align=PP_ALIGN.LEFT, font_name="Aptos", italic=False):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.clear()
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    font = run.font
    font.name = font_name
    font.size = Pt(size)
    font.bold = bold
    font.italic = italic
    font.color.rgb = color
    return box


def add_caption_card(slide, left, top, width, height, title, subtitle):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = WHITE
    shape.line.color.rgb = RGBColor(226, 231, 237)
    shape.line.width = Pt(1)

    add_textbox(slide, left + Inches(0.18), top + Inches(0.12), width - Inches(0.36), Inches(0.28),
                title, 18, NAVY, bold=True)
    add_textbox(slide, left + Inches(0.18), top + Inches(0.46), width - Inches(0.36), Inches(0.62),
                subtitle, 11, SLATE)


def build_slide_one(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, WHITE)

    accent = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(0.22), prs.slide_height
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = NAVY
    accent.line.fill.background()

    soft_panel = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(7.1), Inches(0.55), Inches(5.7), Inches(5.8)
    )
    soft_panel.fill.solid()
    soft_panel.fill.fore_color.rgb = LIGHT
    soft_panel.line.fill.background()

    slide.shapes.add_picture(str(LOGO), Inches(0.9), Inches(0.95), width=Inches(4.1))
    add_textbox(
        slide, Inches(0.92), Inches(2.78), Inches(4.7), Inches(0.7),
        "ADIB Cards Presentation", 24, NAVY, bold=True
    )
    add_textbox(
        slide, Inches(0.94), Inches(3.42), Inches(4.55), Inches(1.1),
        "Simple, visual and readable slides designed around the ADIB card collection.",
        14, SLATE
    )
    add_textbox(
        slide, Inches(0.94), Inches(4.72), Inches(2.0), Inches(0.35),
        "Card portfolio snapshot", 11, CYAN, bold=True
    )
    line = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.94), Inches(5.09), Inches(1.25), Inches(0.04)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = CYAN
    line.line.fill.background()

    card1 = slide.shapes.add_picture(str(CARDS["Titanium Cash Back"]), Inches(7.55), Inches(1.65), width=Inches(3.6))
    card1.rotation = -13
    card2 = slide.shapes.add_picture(str(CARDS["Platinum Cash Back"]), Inches(8.7), Inches(1.45), width=Inches(3.65))
    card2.rotation = 5
    card3 = slide.shapes.add_picture(str(CARDS["Platinum Debit"]), Inches(8.08), Inches(2.4), width=Inches(3.95))
    card3.rotation = -3


def build_slide_two(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, LIGHT)

    add_textbox(slide, Inches(0.72), Inches(0.55), Inches(4.2), Inches(0.6), "Card Portfolio", 26, NAVY, bold=True)
    add_textbox(
        slide, Inches(0.72), Inches(1.12), Inches(7.2), Inches(0.45),
        "Three clean product visuals with short labels for fast reading.",
        13, SLATE
    )

    positions = [Inches(0.72), Inches(4.35), Inches(7.98)]
    captions = [
        ("Titanium Cash Back", "Shari'a-compliant cashback card with a modern silver finish."),
        ("Platinum Cash Back", "Premium cashback option presented with a polished platinum look."),
        ("Platinum Debit", "Contactless platinum debit card designed for everyday convenience."),
    ]

    for idx, (title, subtitle) in enumerate(captions):
        left = positions[idx]
        pic = slide.shapes.add_picture(str(CARDS[title]), left, Inches(1.9), width=Inches(2.95))
        pic.rotation = -2 if idx == 0 else (2 if idx == 1 else 0)
        add_caption_card(slide, left, Inches(4.45), Inches(2.95), Inches(1.25), title, subtitle)


def build_slide_three(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, WHITE)

    banner = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(0.55), Inches(11.9), Inches(1.0)
    )
    banner.fill.solid()
    banner.fill.fore_color.rgb = NAVY
    banner.line.fill.background()

    add_textbox(slide, Inches(1.0), Inches(0.82), Inches(6.3), Inches(0.35), "Featured Card Designs", 24, WHITE, bold=True)
    add_textbox(slide, Inches(1.0), Inches(1.15), Inches(6.3), Inches(0.25), "A balanced comparison between the two highlighted variants.", 12, SKY)

    card_a = slide.shapes.add_picture(str(CARDS["Platinum Cash Back"]), Inches(1.1), Inches(2.0), width=Inches(4.5))
    card_a.rotation = -1
    card_b = slide.shapes.add_picture(str(CARDS["Titanium Cash Back"]), Inches(7.25), Inches(2.0), width=Inches(4.5))
    card_b.rotation = 1

    add_caption_card(
        slide, Inches(1.25), Inches(4.95), Inches(4.1), Inches(0.95),
        "Platinum", "Premium visual finish with a refined, high-contrast layout."
    )
    add_caption_card(
        slide, Inches(7.4), Inches(4.95), Inches(4.1), Inches(0.95),
        "Titanium", "Soft metallic styling that feels modern, simple and professional."
    )


def build_slide_four(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, NAVY)

    band = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.75), Inches(4.5), Inches(4.95)
    )
    band.fill.solid()
    band.fill.fore_color.rgb = RGBColor(11, 80, 162)
    band.line.fill.background()

    add_textbox(slide, Inches(1.18), Inches(1.15), Inches(3.5), Inches(0.5), "Brand Snapshot", 24, WHITE, bold=True)
    add_textbox(
        slide, Inches(1.18), Inches(1.88), Inches(3.3), Inches(2.2),
        "• Clear bilingual branding\n• Strong blue identity\n• Premium card presentation\n• Simple visual structure",
        18, WHITE
    )
    add_textbox(slide, Inches(1.18), Inches(4.5), Inches(3.3), Inches(0.4), "Designed for a clean and readable story.", 12, SKY, italic=True)

    slide.shapes.add_picture(str(LOGO), Inches(6.2), Inches(1.65), width=Inches(5.6))
    slide.shapes.add_picture(str(CARDS["Platinum Debit"]), Inches(7.0), Inches(3.45), width=Inches(3.85))


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "ADIB Simple Visual Presentation"
    prs.core_properties.subject = "ADIB card presentation"
    prs.core_properties.author = "Cursor"

    build_slide_one(prs)
    build_slide_two(prs)
    build_slide_three(prs)
    build_slide_four(prs)

    prs.save(str(OUTPUT_FILE))
    print(f"Created {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
