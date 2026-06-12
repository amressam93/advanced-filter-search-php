#!/usr/bin/env python3
"""Build a clean, visual ADIB Cards PowerPoint deck.

Design goals: simple, visual and readable. Brand navy accents on a light
background, generous whitespace, one idea per slide, large legible type.
"""
import os

from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, "assets")

# Palette
NAVY = RGBColor(0x00, 0x33, 0x66)
NAVY_DEEP = RGBColor(0x00, 0x20, 0x4A)
SKY = RGBColor(0x29, 0xAB, 0xE2)
INK = RGBColor(0x1F, 0x2A, 0x37)
GREY = RGBColor(0x5B, 0x66, 0x72)
LIGHT = RGBColor(0xF4, 0xF6, 0xF9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FONT = "Calibri"
FONT_H = "Calibri"

EMU_IN = 914400
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H
BLANK = prs.slide_layouts[6]


def img_size(path):
    with Image.open(path) as im:
        return im.size


def add_pic_fit(slide, path, left, top, max_w, max_h, align="center", valign="middle"):
    """Place a picture scaled to fit within a box, preserving aspect ratio."""
    w, h = img_size(path)
    ar = w / h
    box_ar = max_w / max_h
    if ar > box_ar:
        new_w = max_w
        new_h = int(max_w / ar)
    else:
        new_h = max_h
        new_w = int(max_h * ar)
    if align == "center":
        x = left + (max_w - new_w) // 2
    elif align == "right":
        x = left + (max_w - new_w)
    else:
        x = left
    if valign == "middle":
        y = top + (max_h - new_h) // 2
    elif valign == "bottom":
        y = top + (max_h - new_h)
    else:
        y = top
    return slide.shapes.add_picture(path, x, y, width=new_w, height=new_h)


def fill_bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def rect(slide, left, top, w, h, color, line=None):
    from pptx.enum.shapes import MSO_SHAPE
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, w, h)
    sp.fill.solid()
    sp.fill.fore_color.rgb = color
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(1)
    sp.shadow.inherit = False
    return sp


def rrect(slide, left, top, w, h, color, line=None):
    from pptx.enum.shapes import MSO_SHAPE
    sp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, w, h)
    sp.fill.solid()
    sp.fill.fore_color.rgb = color
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(1.25)
    sp.shadow.inherit = False
    return sp


def textbox(slide, left, top, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    """lines: list of dicts {text, size, bold, color, font, space_after, space_before}."""
    tb = slide.shapes.add_textbox(left, top, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ln.get("align", align)
        if ln.get("space_after") is not None:
            p.space_after = Pt(ln["space_after"])
        if ln.get("space_before") is not None:
            p.space_before = Pt(ln["space_before"])
        if ln.get("line_spacing"):
            p.line_spacing = ln["line_spacing"]
        run = p.add_run()
        run.text = ln["text"]
        run.font.size = Pt(ln.get("size", 18))
        run.font.bold = ln.get("bold", False)
        run.font.name = ln.get("font", FONT)
        run.font.color.rgb = ln.get("color", INK)
    return tb


def bullets(slide, left, top, w, h, items, size=18, color=INK, gap=10,
            marker_color=SKY):
    tb = slide.shapes.add_textbox(left, top, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        p.line_spacing = 1.05
        r1 = p.add_run()
        r1.text = "\u25B8  "  # small triangle marker
        r1.font.size = Pt(size)
        r1.font.color.rgb = marker_color
        r1.font.bold = True
        r1.font.name = FONT
        head, _, rest = it.partition("|")
        r2 = p.add_run()
        r2.text = head.strip()
        r2.font.size = Pt(size)
        r2.font.bold = bool(rest)
        r2.font.color.rgb = color
        r2.font.name = FONT
        if rest:
            r3 = p.add_run()
            r3.text = "  " + rest.strip()
            r3.font.size = Pt(size)
            r3.font.color.rgb = GREY
            r3.font.name = FONT
    return tb


def section_label(slide, text, left, top, color=SKY):
    textbox(slide, left, top, Inches(5), Inches(0.4),
            [{"text": text.upper(), "size": 14, "bold": True, "color": color}])


def slide_number(slide, n):
    textbox(slide, SLIDE_W - Inches(1.0), SLIDE_H - Inches(0.55),
            Inches(0.7), Inches(0.35),
            [{"text": str(n), "size": 12, "color": GREY, "align": PP_ALIGN.RIGHT}],
            align=PP_ALIGN.RIGHT)


P = os.path.join
LOGO = P(ASSETS, "logo.png")
LOGO_DARK = P(ASSETS, "logo_dark.png")
CARDS3 = P(ASSETS, "cards_three.png")
CARD_PLAT = P(ASSETS, "card_platinum.png")
CARD_TITAN = P(ASSETS, "card_titanium.png")
CARD_CLASSIC = P(ASSETS, "card_classic.png")


# ---------------------------------------------------------------- Slide 1: Title
s = prs.slides.add_slide(BLANK)
fill_bg(s, WHITE)
# left navy accent strip
rect(s, 0, 0, Inches(0.22), SLIDE_H, NAVY)
# logo on the left
add_pic_fit(s, LOGO, Inches(0.85), Inches(0.7), Inches(4.6), Inches(1.4),
            align="left", valign="top")
# title block on the left
textbox(s, Inches(0.9), Inches(2.7), Inches(5.6), Inches(3.2), [
    {"text": "ADIB Cards", "size": 54, "bold": True, "color": NAVY,
     "space_after": 6},
    {"text": "Premium banking, designed around you.", "size": 22,
     "color": GREY, "space_after": 18},
    {"text": "A simple look at our Platinum, Titanium and Classic cards.",
     "size": 16, "color": GREY},
])
# three-cards image on the right
add_pic_fit(s, CARDS3, Inches(6.7), Inches(1.1), Inches(6.2), Inches(5.3),
            align="center", valign="middle")


# ----------------------------------------------------- Slide 2: Card family
s = prs.slides.add_slide(BLANK)
fill_bg(s, LIGHT)
section_label(s, "Overview", Inches(0.9), Inches(0.6))
textbox(s, Inches(0.9), Inches(0.95), Inches(11), Inches(1.0), [
    {"text": "Our Card Family", "size": 40, "bold": True, "color": NAVY},
])
textbox(s, Inches(0.9), Inches(1.85), Inches(11.5), Inches(0.7), [
    {"text": "Three tiers, one promise \u2013 secure, contactless and built for "
             "everyday life.", "size": 18, "color": GREY}])

cards = [(CARD_CLASSIC, "Classic", "Everyday essentials"),
         (CARD_TITAN, "Titanium", "More rewards & travel"),
         (CARD_PLAT, "Platinum", "Premium privileges")]
col_w = Inches(3.7)
gap = Inches(0.35)
start_x = Inches(0.9)
y = Inches(2.9)
for i, (img, title, sub) in enumerate(cards):
    x = start_x + i * (col_w + gap)
    panel = rrect(s, x, y, col_w, Inches(3.4), WHITE)
    add_pic_fit(s, img, x + Inches(0.25), y + Inches(0.3),
                col_w - Inches(0.5), Inches(2.0))
    textbox(s, x + Inches(0.25), y + Inches(2.45), col_w - Inches(0.5), Inches(0.9), [
        {"text": title, "size": 22, "bold": True, "color": NAVY,
         "align": PP_ALIGN.CENTER, "space_after": 2},
        {"text": sub, "size": 14, "color": GREY, "align": PP_ALIGN.CENTER},
    ], align=PP_ALIGN.CENTER)
slide_number(s, 2)


# ------------------------------------------------ Card detail slides helper
def card_detail(n, img, label, title, tagline, feats, image_left=True):
    s = prs.slides.add_slide(BLANK)
    fill_bg(s, WHITE)
    rect(s, 0, 0, SLIDE_W, Inches(0.16), SKY)
    img_w = Inches(5.9)
    img_h = Inches(4.4)
    if image_left:
        ix, iy = Inches(0.7), Inches(1.7)
        tx = Inches(7.0)
    else:
        ix, iy = Inches(6.7), Inches(1.7)
        tx = Inches(0.9)
    add_pic_fit(s, img, ix, iy, img_w, img_h, align="center", valign="middle")
    section_label(s, label, tx, Inches(1.0))
    textbox(s, tx, Inches(1.35), Inches(5.6), Inches(1.0), [
        {"text": title, "size": 38, "bold": True, "color": NAVY}])
    textbox(s, tx, Inches(2.25), Inches(5.6), Inches(0.7), [
        {"text": tagline, "size": 18, "color": SKY, "bold": True}])
    bullets(s, tx, Inches(3.1), Inches(5.7), Inches(3.3), feats, size=17, gap=12)
    slide_number(s, n)
    return s


card_detail(3, CARD_PLAT, "Platinum", "Platinum Card",
            "Premium privileges, worldwide.", [
                "Airport lounge access | at major global hubs",
                "Higher limits | for bigger plans and purchases",
                "Travel & purchase protection | included",
                "Concierge support | available 24/7",
            ], image_left=True)

card_detail(4, CARD_TITAN, "Titanium", "Titanium Card",
            "More rewards on the go.", [
                "Accelerated rewards | on everyday spend",
                "Global acceptance | wherever Mastercard is welcome",
                "Contactless & secure | tap to pay in seconds",
                "Smart spend controls | in the ADIB app",
            ], image_left=False)

card_detail(5, CARD_CLASSIC, "Classic", "Classic Card",
            "Your everyday essentials.", [
                "Simple, low-cost | banking for daily life",
                "Secure chip & PIN | with EMV protection",
                "Contactless payments | fast and easy",
                "Full control | manage anytime in the app",
            ], image_left=True)


# ----------------------------------------------------- Slide 6: Key benefits
s = prs.slides.add_slide(BLANK)
fill_bg(s, LIGHT)
section_label(s, "Why ADIB Cards", Inches(0.9), Inches(0.6))
textbox(s, Inches(0.9), Inches(0.95), Inches(11), Inches(1.0), [
    {"text": "Built Around Four Simple Promises", "size": 36, "bold": True,
     "color": NAVY}])
benefits = [
    ("Secure", "EMV chip, PIN and real-time fraud monitoring keep you safe."),
    ("Contactless", "Tap to pay in seconds, in stores and online."),
    ("Rewarding", "Earn value on the spending you already do."),
    ("Sharia-compliant", "Banking that aligns with your values."),
]
bw = Inches(5.65)
bh = Inches(1.95)
gx = Inches(0.9)
gy = Inches(2.3)
for i, (t, d) in enumerate(benefits):
    r = i // 2
    c = i % 2
    x = gx + c * (bw + Inches(0.4))
    y = gy + r * (bh + Inches(0.35))
    rrect(s, x, y, bw, bh, WHITE)
    rect(s, x, y, Inches(0.14), bh, SKY)
    textbox(s, x + Inches(0.45), y + Inches(0.28), bw - Inches(0.7), bh - Inches(0.5), [
        {"text": t, "size": 22, "bold": True, "color": NAVY, "space_after": 6},
        {"text": d, "size": 15, "color": GREY, "line_spacing": 1.05},
    ])
slide_number(s, 6)


# ----------------------------------------------------- Slide 7: Comparison
s = prs.slides.add_slide(BLANK)
fill_bg(s, WHITE)
rect(s, 0, 0, SLIDE_W, Inches(0.16), SKY)
section_label(s, "At a Glance", Inches(0.9), Inches(0.55))
textbox(s, Inches(0.9), Inches(0.9), Inches(11), Inches(0.9), [
    {"text": "Compare the Cards", "size": 36, "bold": True, "color": NAVY}])

rows = [
    ["Feature", "Classic", "Titanium", "Platinum"],
    ["Contactless payments", "Yes", "Yes", "Yes"],
    ["Rewards rate", "Standard", "Higher", "Highest"],
    ["Travel benefits", "\u2013", "Selected", "Extensive"],
    ["Airport lounge access", "\u2013", "\u2013", "Included"],
    ["Concierge service", "\u2013", "\u2013", "24/7"],
]
tbl_left = Inches(0.9)
tbl_top = Inches(2.05)
tbl_w = Inches(11.5)
tbl_h = Inches(4.6)
shape = s.shapes.add_table(len(rows), 4, tbl_left, tbl_top, tbl_w, tbl_h)
table = shape.table
table.columns[0].width = Inches(4.3)
for c in range(1, 4):
    table.columns[c].width = Inches((11.5 - 4.3) / 3)
for r, row in enumerate(rows):
    table.rows[r].height = Inches(tbl_h.inches / len(rows))
    for c, val in enumerate(row):
        cell = table.cell(r, c)
        cell.margin_left = Inches(0.18)
        cell.margin_right = Inches(0.1)
        cell.margin_top = Inches(0.06)
        cell.margin_bottom = Inches(0.06)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER
        run = p.add_run()
        run.text = val
        run.font.name = FONT
        if r == 0:
            run.font.size = Pt(18)
            run.font.bold = True
            run.font.color.rgb = WHITE
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY
        else:
            run.font.size = Pt(16)
            run.font.bold = (c == 0)
            run.font.color.rgb = INK if c == 0 else GREY
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if r % 2 else LIGHT
slide_number(s, 7)


# ----------------------------------------------------- Slide 8: Closing
s = prs.slides.add_slide(BLANK)
fill_bg(s, NAVY_DEEP)
rect(s, 0, 0, SLIDE_W, SLIDE_H, NAVY_DEEP)
add_pic_fit(s, LOGO_DARK, Inches(4.0), Inches(1.7), Inches(5.3), Inches(1.7),
            align="center", valign="middle")
textbox(s, Inches(1.5), Inches(3.9), Inches(10.3), Inches(1.6), [
    {"text": "Thank You", "size": 46, "bold": True, "color": WHITE,
     "align": PP_ALIGN.CENTER, "space_after": 8},
    {"text": "Choose the ADIB card that fits your life.", "size": 20,
     "color": SKY, "align": PP_ALIGN.CENTER},
], align=PP_ALIGN.CENTER)
textbox(s, Inches(1.5), Inches(6.2), Inches(10.3), Inches(0.6), [
    {"text": "www.adib.ae   |   Abu Dhabi Islamic Bank", "size": 15,
     "color": RGBColor(0xB8, 0xC6, 0xD8), "align": PP_ALIGN.CENTER}],
    align=PP_ALIGN.CENTER)


OUT = os.path.join(HERE, "ADIB_Cards_Presentation.pptx")
prs.save(OUT)
print("Saved", OUT, "with", len(prs.slides._sldIdLst), "slides")
