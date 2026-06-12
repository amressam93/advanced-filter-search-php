from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


BASE_DIR = Path(__file__).resolve().parent
ASSETS_DIR = BASE_DIR / "assets"
OUTPUT_PATH = BASE_DIR / "adib_simple_visual_presentation.pptx"


WHITE = RGBColor(255, 255, 255)
NAVY = RGBColor(0, 66, 138)
DEEP_NAVY = RGBColor(9, 33, 79)
SKY = RGBColor(53, 181, 238)
PALE_BLUE = RGBColor(230, 244, 255)
LIGHT_BG = RGBColor(247, 250, 252)
MUTED = RGBColor(86, 99, 115)
TEXT_DARK = RGBColor(31, 41, 55)
SHADOW = RGBColor(21, 32, 43)


def set_background(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_shape(slide, shape_type, left, top, width, height, fill_color, line_color=None, transparency=0):
    shape = slide.shapes.add_shape(shape_type, left, top, width, height)
    fill = shape.fill
    fill.solid()
    fill.fore_color.rgb = fill_color
    fill.transparency = transparency

    if line_color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line_color

    return shape


def add_textbox(
    slide,
    text,
    left,
    top,
    width,
    height,
    font_size,
    color=TEXT_DARK,
    bold=False,
    align=PP_ALIGN.LEFT,
    font_name="Aptos",
):
    textbox = slide.shapes.add_textbox(left, top, width, height)
    frame = textbox.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE

    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    return textbox


def add_bullets(slide, items, left, top, width, height):
    textbox = slide.shapes.add_textbox(left, top, width, height)
    frame = textbox.text_frame
    frame.clear()
    frame.word_wrap = True

    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = f"• {item}"
        paragraph.level = 0
        paragraph.alignment = PP_ALIGN.LEFT
        paragraph.space_after = Pt(10)
        for run in paragraph.runs:
            run.font.name = "Aptos"
            run.font.size = Pt(22)
            run.font.color.rgb = TEXT_DARK

    return textbox


def add_picture_with_shadow(slide, image_path, left, top, width=None, height=None, rotation=0, shadow_dx=0.07, shadow_dy=0.07):
    shadow_left = left + Inches(shadow_dx)
    shadow_top = top + Inches(shadow_dy)
    shadow_width = width if width is not None else Inches(1)
    shadow_height = height if height is not None else Inches(1)

    shadow = add_shape(
        slide,
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        shadow_left,
        shadow_top,
        shadow_width,
        shadow_height,
        SHADOW,
        transparency=0.82,
    )
    shadow.rotation = rotation

    picture = slide.shapes.add_picture(str(image_path), left, top, width=width, height=height)
    picture.rotation = rotation
    return picture


def add_title(slide, title, subtitle=None):
    add_textbox(slide, title, Inches(0.9), Inches(0.55), Inches(7.2), Inches(0.8), 28, NAVY, True)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.92), Inches(1.18), Inches(0.9), Inches(0.05), SKY)
    if subtitle:
        add_textbox(slide, subtitle, Inches(0.9), Inches(1.28), Inches(8.8), Inches(0.6), 13, MUTED)


def add_footer(slide):
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.9), Inches(7.0), Inches(11.4), Inches(0.02), PALE_BLUE)
    add_textbox(
        slide,
        "ADIB | Simple visual deck",
        Inches(9.0),
        Inches(7.05),
        Inches(3.0),
        Inches(0.25),
        10,
        MUTED,
        align=PP_ALIGN.RIGHT,
    )


def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank = prs.slide_layouts[6]

    logo = ASSETS_DIR / "logo_full.png"
    card_cashback = ASSETS_DIR / "card_cashback.png"
    card_platinum = ASSETS_DIR / "card_platinum.png"
    card_business = ASSETS_DIR / "card_business_platinum.png"

    # Slide 1: cover
    slide = prs.slides.add_slide(blank)
    set_background(slide, LIGHT_BG)

    add_shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, Inches(8.0), Inches(0.7), Inches(4.2), Inches(4.2), PALE_BLUE, transparency=0.1)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, Inches(9.25), Inches(2.0), Inches(2.2), Inches(2.2), WHITE, transparency=0.0)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.9), Inches(0.55), Inches(0.15), Inches(0.7), SKY)

    slide.shapes.add_picture(str(logo), Inches(1.1), Inches(1.55), width=Inches(5.0))
    add_textbox(slide, "Simple visual presentation", Inches(1.15), Inches(3.15), Inches(4.8), Inches(0.55), 22, NAVY, True)
    add_textbox(
        slide,
        "Clean layout, readable text, and a premium card-focused first slide.",
        Inches(1.15),
        Inches(3.7),
        Inches(4.9),
        Inches(0.8),
        16,
        MUTED,
    )
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(1.15), Inches(4.75), Inches(2.0), Inches(0.42), PALE_BLUE)
    add_textbox(slide, "Cover concept", Inches(1.2), Inches(4.78), Inches(1.9), Inches(0.34), 12, NAVY, True, PP_ALIGN.CENTER)

    add_picture_with_shadow(slide, card_business, Inches(7.25), Inches(2.1), height=Inches(2.55), rotation=-17)
    add_picture_with_shadow(slide, card_cashback, Inches(8.2), Inches(1.9), height=Inches(2.75), rotation=-2)
    add_picture_with_shadow(slide, card_platinum, Inches(9.15), Inches(2.05), height=Inches(2.55), rotation=14)

    add_footer(slide)

    # Slide 2: card gallery
    slide = prs.slides.add_slide(blank)
    set_background(slide, WHITE)
    add_title(slide, "Card gallery", "A simple three-card layout that keeps the visuals clear and balanced.")

    card_specs = [
        (card_cashback, "Cashback Visa Covered Card", Inches(0.9)),
        (card_platinum, "Exceed Visa Platinum Covered Card", Inches(4.55)),
        (card_business, "Business Visa Platinum Covered Card", Inches(8.2)),
    ]

    for image_path, label, left in card_specs:
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left + Inches(0.08), Inches(1.98), Inches(3.15), Inches(3.7), SHADOW, transparency=0.9)
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, Inches(1.9), Inches(3.15), Inches(3.7), LIGHT_BG, line_color=PALE_BLUE)
        slide.shapes.add_picture(str(image_path), left + Inches(0.2), Inches(2.18), width=Inches(2.75))
        add_textbox(slide, label, left + Inches(0.22), Inches(5.05), Inches(2.7), Inches(0.48), 14, NAVY, True, PP_ALIGN.CENTER)
        add_textbox(slide, "Simple caption area for editing.", left + Inches(0.22), Inches(5.4), Inches(2.7), Inches(0.36), 11, MUTED, False, PP_ALIGN.CENTER)

    add_footer(slide)

    # Slide 3: design principles
    slide = prs.slides.add_slide(blank)
    set_background(slide, LIGHT_BG)
    add_title(slide, "Design principles", "Built to stay visual, readable, and easy to update.")

    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(1.8), Inches(5.25), Inches(4.4), WHITE, line_color=PALE_BLUE)
    add_textbox(slide, "Keep each slide focused", Inches(1.2), Inches(2.05), Inches(3.6), Inches(0.4), 18, NAVY, True)
    add_bullets(
        slide,
        [
            "Use one main message per slide.",
            "Limit text to short phrases and labels.",
            "Let the card images carry most of the visual weight.",
            "Use the blue ADIB palette for a clean, consistent look.",
        ],
        Inches(1.2),
        Inches(2.5),
        Inches(4.5),
        Inches(2.9),
    )

    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(6.55), Inches(1.8), Inches(5.85), Inches(4.4), WHITE, line_color=PALE_BLUE)
    add_textbox(slide, "Recommended visual balance", Inches(6.9), Inches(2.05), Inches(3.6), Inches(0.4), 18, NAVY, True)
    add_textbox(slide, "Large image area", Inches(7.0), Inches(2.6), Inches(2.2), Inches(0.35), 12, NAVY, True, PP_ALIGN.CENTER)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(6.95), Inches(2.95), Inches(4.9), Inches(2.45), PALE_BLUE, transparency=0.18)
    add_picture_with_shadow(slide, card_platinum, Inches(7.55), Inches(3.08), width=Inches(3.7), rotation=-6, shadow_dx=0.05, shadow_dy=0.05)
    add_textbox(slide, "Small text block", Inches(9.1), Inches(5.55), Inches(2.0), Inches(0.3), 12, MUTED)

    add_footer(slide)

    # Slide 4: closing
    slide = prs.slides.add_slide(blank)
    set_background(slide, WHITE)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, Inches(9.5), Inches(0.75), Inches(2.2), Inches(2.2), PALE_BLUE, transparency=0.2)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.9), Inches(0.55), Inches(0.15), Inches(0.7), SKY)
    add_textbox(slide, "Thank you", Inches(1.1), Inches(1.65), Inches(3.5), Inches(0.8), 30, NAVY, True)
    add_textbox(
        slide,
        "This deck is set up as a simple visual presentation and can be edited easily in PowerPoint.",
        Inches(1.1),
        Inches(2.35),
        Inches(5.3),
        Inches(0.8),
        17,
        MUTED,
    )
    slide.shapes.add_picture(str(logo), Inches(1.1), Inches(3.45), width=Inches(3.8))
    add_picture_with_shadow(slide, card_cashback, Inches(8.5), Inches(2.0), height=Inches(2.6), rotation=-10)
    add_picture_with_shadow(slide, card_business, Inches(9.55), Inches(2.25), height=Inches(2.45), rotation=10)
    add_footer(slide)

    prs.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    path = create_presentation()
    print(f"Created presentation: {path}")
