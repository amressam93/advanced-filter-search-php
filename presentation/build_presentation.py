"""
Build a clean, visual ADIB Cards PowerPoint presentation.

Run:  python3 build_presentation.py
Output: ADIB_Cards_Presentation.pptx (next to this script)

Images used (in ./assets):
  - logo_white.png      ADIB logo (white, for blue panels)
  - cards_fan.png       Three cards fanned (Blue / Titanium / Platinum)
  - platinum_card.png   ADIB Platinum Cash Back card
  - titanium_card.png   ADIB Titanium Cash Back card
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, "assets")

# ---- Brand palette -------------------------------------------------------
DEEP   = RGBColor(0x00, 0x33, 0x60)   # deep navy
BLUE   = RGBColor(0x00, 0x5E, 0xA8)   # ADIB blue
ACCENT = RGBColor(0x12, 0xA5, 0xE3)   # light accent blue
SKY    = RGBColor(0xE9, 0xF3, 0xFB)   # very light blue
LIGHT  = RGBColor(0xF4, 0xF7, 0xFB)   # near-white panel
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
INK    = RGBColor(0x1B, 0x2A, 0x3A)   # heading text
GREY   = RGBColor(0x53, 0x63, 0x72)   # body text
GOLD   = RGBColor(0xF2, 0x9A, 0x1F)   # mastercard amber accent

FONT_H = "Segoe UI Semibold"
FONT_B = "Segoe UI"

EMU_IN = 914400

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


# ---- Helpers -------------------------------------------------------------
def slide():
    return prs.slides.add_slide(BLANK)


def bg(s, color):
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = color


def rect(s, x, y, w, h, color, shape=MSO_SHAPE.RECTANGLE, line=None, line_w=1.0):
    sp = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.fill.solid()
    sp.fill.fore_color.rgb = color
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(line_w)
    sp.shadow.inherit = False
    return sp


def text(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         space_after=6, line_spacing=1.0):
    """runs: list of paragraphs; each paragraph is list of (txt, size, color, bold, font)."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        p.space_before = Pt(0)
        p.line_spacing = line_spacing
        for (txt, size, color, bold, font) in para:
            r = p.add_run()
            r.text = txt
            r.font.size = Pt(size)
            r.font.color.rgb = color
            r.font.bold = bold
            r.font.name = font
    return tb


def pic_fit(s, path, x, y, w, h, align="center", valign="middle"):
    """Place picture inside box (x,y,w,h) in inches, preserving aspect ratio."""
    iw, ih = Image.open(path).size
    box_r = w / h
    img_r = iw / ih
    if img_r > box_r:
        nw = w
        nh = w / img_r
    else:
        nh = h
        nw = h * img_r
    if align == "center":
        px = x + (w - nw) / 2
    elif align == "left":
        px = x
    else:
        px = x + (w - nw)
    if valign == "middle":
        py = y + (h - nh) / 2
    elif valign == "top":
        py = y
    else:
        py = y + (h - nh)
    return s.shapes.add_picture(path, Inches(px), Inches(py), Inches(nw), Inches(nh))


def chip(s, x, y, w, h, fill, title, sub, t_color=WHITE, s_color=None):
    c = rect(s, x, y, w, h, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    _round(c, 0.12)
    if s_color is None:
        s_color = t_color
    text(s, x + 0.25, y, w - 0.5, h,
         [[(title, 20, t_color, True, FONT_H)],
          [(sub, 12.5, s_color, False, FONT_B)]],
         anchor=MSO_ANCHOR.MIDDLE, space_after=3, line_spacing=1.0)
    return c


def _round(shape, val=0.1):
    try:
        shape.adjustments[0] = val
    except Exception:
        pass


def num_badge(s, x, y, n, color=ACCENT):
    b = rect(s, x, y, 0.5, 0.5, color, shape=MSO_SHAPE.OVAL)
    text(s, x, y, 0.5, 0.5, [[(str(n), 18, WHITE, True, FONT_H)]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, space_after=0)


def footer(s, dark=False):
    col = WHITE if dark else GREY
    text(s, 0.55, SH/EMU_IN - 0.5, 6, 0.3,
         [[("ADIB  |  Banking as it should be", 10, col, False, FONT_B)]],
         space_after=0)


LOGO = os.path.join(ASSETS, "logo_white.png")
FAN = os.path.join(ASSETS, "cards_fan.png")
PLAT = os.path.join(ASSETS, "platinum_card.png")
TITA = os.path.join(ASSETS, "titanium_card.png")


# =========================================================================
# SLIDE 1 — Title : logo on the LEFT, three cards on the RIGHT
# =========================================================================
s = slide()
bg(s, WHITE)
# left blue panel
rect(s, 0, 0, 5.5, 7.5, DEEP)
rect(s, 5.5, 0, 0.12, 7.5, ACCENT)        # thin accent divider
# logo on left panel
pic_fit(s, LOGO, 0.7, 0.85, 3.9, 1.0, align="left")
# accent underline
rect(s, 0.75, 2.15, 1.5, 0.06, ACCENT)
# title text on left
text(s, 0.75, 2.65, 4.4, 3.0,
     [[("ADIB Cards", 44, WHITE, True, FONT_H)],
      [("Smart, Sharia-compliant cards", 19, SKY, False, FONT_B)],
      [("for the way you live and spend.", 19, SKY, False, FONT_B)]],
     space_after=6, line_spacing=1.05)
text(s, 0.75, 6.35, 4.4, 0.6,
     [[("Banking as it should be", 14, ACCENT, True, FONT_B)]])
# three cards on the right (white area)
pic_fit(s, FAN, 5.85, 1.3, 7.1, 5.0, align="center")
text(s, 5.85, 6.35, 7.0, 0.5,
     [[("Blue   •   Titanium   •   Platinum", 14, GREY, True, FONT_B)]],
     align=PP_ALIGN.CENTER)


# =========================================================================
# SLIDE 2 — Overview
# =========================================================================
s = slide()
bg(s, LIGHT)
rect(s, 0, 0, 13.333, 1.45, DEEP)
pic_fit(s, LOGO, 10.0, 0.35, 2.8, 0.75, align="right")
text(s, 0.55, 0, 8, 1.45,
     [[("Why ADIB Cards", 28, WHITE, True, FONT_H)]],
     anchor=MSO_ANCHOR.MIDDLE)

text(s, 0.55, 1.85, 12.2, 1.1,
     [[("ADIB cards are designed to be simple, rewarding and secure — "
        "giving you real cash back on everyday spending while staying fully "
        "Sharia-compliant.", 18, GREY, False, FONT_B)]],
     line_spacing=1.15)

cards = [
    (SKY,  "Cash Back", "Earn on every local purchase, with no minimum or maximum limit.", DEEP, GREY),
    (BLUE, "Sharia-Compliant", "MasterCard cards issued in line with Islamic principles.", WHITE, SKY),
    (DEEP, "Secure & Global", "Contactless chip technology, accepted worldwide.", WHITE, SKY),
]
cx = 0.55
cw = 3.95
gap = 0.3
for fill, t, sub, tc, sc in cards:
    c = rect(s, cx, 3.25, cw, 2.9, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    _round(c, 0.06)
    rect(s, cx + 0.45, 3.7, 0.7, 0.12, ACCENT if fill != BLUE else WHITE)
    text(s, cx + 0.45, 4.05, cw - 0.9, 2.0,
         [[(t, 22, tc, True, FONT_H)],
          [("", 6, tc, False, FONT_B)],
          [(sub, 14.5, sc, False, FONT_B)]],
         line_spacing=1.12, space_after=4)
    cx += cw + gap
footer(s)


# =========================================================================
# SLIDE 3 — Platinum card
# =========================================================================
s = slide()
bg(s, WHITE)
rect(s, 0, 0, 6.4, 7.5, SKY)
pic_fit(s, PLAT, 0.2, 1.4, 6.0, 4.7, align="center")
text(s, 0.6, 0.7, 5.4, 0.6, [[("ADIB Platinum", 16, BLUE, True, FONT_B)]])

text(s, 7.0, 0.95, 5.9, 1.6,
     [[("Platinum", 40, INK, True, FONT_H)],
      [("Cash Back Card", 26, BLUE, True, FONT_H)]],
     space_after=2, line_spacing=1.0)
rect(s, 7.05, 2.55, 1.5, 0.06, ACCENT)

feats = [
    "Generous cash back on all your local spending.",
    "No minimum or maximum cash-back limits.",
    "Contactless payments for fast, secure checkout.",
    "Worldwide acceptance on the MasterCard network.",
    "Optional Takaful cover for added peace of mind.",
]
fy = 2.95
for i, f in enumerate(feats, 1):
    num_badge(s, 7.05, fy, i)
    text(s, 7.75, fy - 0.02, 5.2, 0.6,
         [[(f, 15.5, GREY, False, FONT_B)]], anchor=MSO_ANCHOR.MIDDLE)
    fy += 0.82
footer(s)


# =========================================================================
# SLIDE 4 — Titanium card (mirrored layout)
# =========================================================================
s = slide()
bg(s, WHITE)
rect(s, 6.93, 0, 6.4, 7.5, RGBColor(0x1C, 0x24, 0x2E))
pic_fit(s, TITA, 7.1, 1.4, 6.0, 4.7, align="center")
text(s, 7.5, 0.7, 5.4, 0.6, [[("ADIB Titanium", 16, ACCENT, True, FONT_B)]])

text(s, 0.6, 0.95, 6.0, 1.6,
     [[("Titanium", 40, INK, True, FONT_H)],
      [("Cash Back Card", 26, BLUE, True, FONT_H)]],
     space_after=2, line_spacing=1.0)
rect(s, 0.65, 2.55, 1.5, 0.06, ACCENT)

feats = [
    "Premium card with cash back on every purchase.",
    "Spend locally and globally with total freedom.",
    "Latest chip & contactless security technology.",
    "Sharia-compliant MasterCard you can trust.",
    "Flexible installment plans of 3 / 6 / 9 / 12 months.",
]
fy = 2.95
for i, f in enumerate(feats, 1):
    num_badge(s, 0.65, fy, i)
    text(s, 1.35, fy - 0.02, 5.2, 0.6,
         [[(f, 15.5, GREY, False, FONT_B)]], anchor=MSO_ANCHOR.MIDDLE)
    fy += 0.82
footer(s)


# =========================================================================
# SLIDE 5 — Comparison
# =========================================================================
s = slide()
bg(s, LIGHT)
rect(s, 0, 0, 13.333, 1.45, DEEP)
pic_fit(s, LOGO, 10.0, 0.35, 2.8, 0.75, align="right")
text(s, 0.55, 0, 8, 1.45, [[("Choose Your Card", 28, WHITE, True, FONT_H)]],
     anchor=MSO_ANCHOR.MIDDLE)

rows = [
    ("Feature", "Platinum", "Titanium", True),
    ("Cash back on local spend", "Yes", "Yes", False),
    ("No min / max limits", "Yes", "Yes", False),
    ("Contactless payment", "Yes", "Yes", False),
    ("Global MasterCard acceptance", "Yes", "Yes", False),
    ("Installment plans (3-12 months)", "Yes", "Yes", False),
    ("Takaful protection (optional)", "Yes", "Yes", False),
]
tx, ty, tw = 0.85, 1.9, 11.6
c1 = tw * 0.5
c2 = tw * 0.25
rh = 0.66
y = ty
for i, (a, b, c, head) in enumerate(rows):
    fill = DEEP if head else (WHITE if i % 2 else SKY)
    rect(s, tx, y, tw, rh, fill, shape=MSO_SHAPE.RECTANGLE)
    ac = WHITE if head else INK
    bc = WHITE if head else BLUE
    text(s, tx + 0.3, y, c1 - 0.3, rh, [[(a, 15, ac, head, FONT_B)]],
         anchor=MSO_ANCHOR.MIDDLE)
    text(s, tx + c1, y, c2, rh, [[(b, 15, bc, True, FONT_B)]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, tx + c1 + c2, y, c2, rh, [[(c, 15, bc, True, FONT_B)]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    y += rh
footer(s)


# =========================================================================
# SLIDE 6 — Key benefits grid
# =========================================================================
s = slide()
bg(s, WHITE)
text(s, 0.6, 0.65, 12, 1.0, [[("Benefits at a Glance", 30, INK, True, FONT_H)]])
rect(s, 0.65, 1.55, 1.6, 0.06, ACCENT)

bens = [
    ("Real Cash Back", "Get rewarded on every local purchase you make."),
    ("100% Sharia-Compliant", "Cards structured in line with Islamic principles."),
    ("Smart Security", "Chip & contactless protection on every payment."),
    ("Pay in Installments", "Split purchases over 3, 6, 9 or 12 months."),
    ("Global Acceptance", "Use your card anywhere MasterCard is welcome."),
    ("Takaful Cover", "Optional protection for total peace of mind."),
]
gx0, gy0 = 0.65, 2.0
gw, gh = 3.95, 1.95
gapx, gapy = 0.28, 0.28
for idx, (t, sub) in enumerate(bens):
    r = idx // 3
    col = idx % 3
    x = gx0 + col * (gw + gapx)
    y = gy0 + r * (gh + gapy)
    card = rect(s, x, y, gw, gh, LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    _round(card, 0.07)
    card.line.color.rgb = SKY
    card.line.width = Pt(1)
    rect(s, x, y, 0.14, gh, ACCENT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, x + 0.4, y + 0.28, gw - 0.7, gh - 0.5,
         [[(t, 19, DEEP, True, FONT_H)],
          [("", 5, DEEP, False, FONT_B)],
          [(sub, 13.5, GREY, False, FONT_B)]],
         line_spacing=1.1, space_after=3)
footer(s)


# =========================================================================
# SLIDE 7 — Closing / contact
# =========================================================================
s = slide()
bg(s, DEEP)
rect(s, 0, 0, 13.333, 0.18, ACCENT)
pic_fit(s, LOGO, 4.67, 1.5, 4.0, 1.1, align="center")
text(s, 1, 3.0, 11.333, 1.2,
     [[("Apply for your ADIB card today", 34, WHITE, True, FONT_H)]],
     align=PP_ALIGN.CENTER)
text(s, 1, 4.1, 11.333, 0.8,
     [[("Simple. Rewarding. Sharia-compliant.", 18, SKY, False, FONT_B)]],
     align=PP_ALIGN.CENTER)
# contact chips
items = ["adib.ae", "600 543216", "ADIB Mobile App"]
cw2 = 3.4
total = cw2 * 3 + 0.6
sx = (13.333 - total) / 2
for it in items:
    c = rect(s, sx, 5.25, cw2, 0.85, BLUE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    _round(c, 0.5)
    text(s, sx, 5.25, cw2, 0.85, [[(it, 16, WHITE, True, FONT_B)]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    sx += cw2 + 0.3
text(s, 0, 6.85, 13.333, 0.4,
     [[("Banking as it should be", 13, ACCENT, True, FONT_B)]],
     align=PP_ALIGN.CENTER)


out = os.path.join(HERE, "ADIB_Cards_Presentation.pptx")
prs.save(out)
print("Saved:", out, "(%d slides)" % len(prs.slides._sldIdLst))
